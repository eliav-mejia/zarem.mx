#!/usr/bin/env python3
"""Rebuild colegio_la_florida/db/productos.json from the Odoo master catalogue (Catalogo.xlsx), and write each row's
photo URL into the catalogue's `imagen` column.

    ~/.venvs/xl/bin/python tools/importar-catalogo.py [ruta/a/Catalogo.xlsx]

Catalogo.xlsx has one row per variant (folio = Odoo default_code, e.g. CLF02V08). Rows are grouped into one web
product per (productos, sexo): the polo diario for niño and niña are two products. Each product keeps its sizes in
catalogue order with that size's folio and price (`variants`); `price_mxn` is the lowest one ("Desde").
Web-only fields (slug, line, type, texts) live in PRODUCTS below.
Photos are named after the Odoo product: img/producto/<colegio>-<productos>-<sexo>.webp, e.g.
colegio-la-florida-playera-polo-diario-nino.webp (see photo_for). Without that file the product uses its PRODUCTS
`image` (a framed crop of a section photo) or sin-foto. Then run tools/generar-paginas-producto.py.
Needs openpyxl.
"""
import json
import re
import unicodedata
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DB = ROOT / 'colegio_la_florida' / 'db' / 'productos.json'
DEFAULT_XLSX = ROOT.parent.parent / 'builds' / 'Odoo' / 'Odoo_catalogues' / 'zarem.mx' / 'Catalogo.xlsx'
SITE = 'https://zarem.mx'
NO_PHOTO = '/img/producto/sin-foto.webp'
COLEGIO = 'Colegio La Florida'

SIZE = {'Chica': 'CH', 'Mediana': 'M', 'Grande': 'G', 'Extra Grande': 'EG'}
FIT = {'Niña': 'nina', 'Niño': 'nino', 'Unisex': 'unisex'}

# (productos, sexo) in Catalogo.xlsx → web product. Texts marked by the client as pending stay empty (the page hides them).
PRODUCTS = {
    ('Playera polo diario', 'Niño'): dict(
        id='cl-102', slug='polo-diario-nino', line_id='diario', type_id='playeras', name='Polo diario niño',
        short='Piqué blanco con vivos y escudo bordado.',
        description='Polo de piqué blanco con cuello y puños tejidos con vivos, y el escudo del colegio bordado en el pecho. Corte para niño.',
        fabric='Piqué 100 % algodón peinado, 220 g/m²', care='Lavar a 30 °C del revés; no usar blanqueador sobre el escudo.',
        features=['Escudo bordado, no estampado', 'Cuello que no se enrosca', 'Vivos tejidos en cuello y manga'], featured=True),
    ('Playera polo diario', 'Niña'): dict(
        id='cl-103', slug='polo-diario-nina', line_id='diario', type_id='playeras', name='Polo diario niña',
        short='Blanca, cuello tejido y escudo bordado.',
        description='Playera polo blanca de diario con el escudo del colegio bordado. Corte para niña; se usa bajo el jumper y el suéter.',
        featured=True),
    ('Playera polo deportes', 'Niño'): dict(
        id='cl-104', slug='polo-deportes-nino', line_id='deportes', type_id='playeras', name='Polo deportes niño',
        short='Playera polo para educación física.',
        description='Playera polo del uniforme de deportes. Corte para niño.'),
    ('Playera polo deportes', 'Niña'): dict(
        id='cl-105', slug='polo-deportes-nina', line_id='deportes', type_id='playeras', name='Polo deportes niña',
        short='Playera polo para educación física.',
        description='Playera polo del uniforme de deportes. Corte para niña.'),
    ('Sueter abierto', 'Niña'): dict(
        id='cl-106', slug='sueter-abierto-nina', line_id='diario', type_id='sueteres', name='Suéter abierto niña',
        short='Cárdigan café, tejido de trenzas.',
        description='Suéter abierto (cárdigan) de punto café con tejido de trenzas, cuello en V y botones al frente.',
        fabric='Acrílico antibolitas 70 % · algodón 30 %', care='Lavar en frío a mano o ciclo delicado; secar extendido.',
        features=['Antibolitas', 'Puños que no se deforman', 'Botones cosidos en cruz'], featured=True),
    ('Jumper', 'Niña'): dict(
        id='cl-107', slug='jumper-nina', line_id='diario', type_id='jumpers', name='Jumper',
        short='Gris jaspeado, con tablones.',
        description='Jumper gris jaspeado de tirantes con tablones al frente y dobladillo amplio para alargar conforme crece.',
        fabric='Gabardina jaspeada 65 % poliéster · 35 % rayón', care='Lavar en frío; planchar los tablones del revés.',
        features=['Dobladillo de 5 cm para alargar', 'Tablones cosidos que no se abren']),
    ('Calcetas', 'Niña'): dict(
        id='cl-108', slug='calcetas-nina', line_id='diario', type_id='accesorios', name='Calcetas niña',
        short='Calcetas del uniforme.', embroidery=False,
        description='Calcetas del uniforme diario para niña.'),
    ('Calcetas', 'Niño'): dict(
        id='cl-109', slug='calcetas-nino', line_id='diario', type_id='accesorios', name='Calcetas niño',
        short='Calcetas del uniforme.', embroidery=False,
        description='Calcetas del uniforme diario para niño.'),
    ('Sueter cerrado', 'Niño'): dict(
        id='cl-110', slug='sueter-cerrado-nino', line_id='diario', type_id='sueteres', name='Suéter cerrado niño',
        short='Café, cuello en V.',
        description='Suéter cerrado de punto liso café con cuello en V y resorte tejido en cuello, puños y cintura.',
        fabric='Acrílico antibolitas 70 % · algodón 30 %', care='Lavar en frío; secar extendido.',
        features=['Antibolitas', 'Codos reforzados']),
    ('Chaleco', 'Niño'): dict(
        id='cl-111', slug='chaleco-nino', line_id='diario', type_id='chalecos', name='Chaleco niño',
        short='Chaleco del uniforme diario.',
        description='Chaleco del uniforme diario para niño.'),
    ('Pantalón de diario', 'Niño'): dict(
        id='cl-112', slug='pantalon-diario-nino', line_id='diario', type_id='pantalones', name='Pantalón de diario niño',
        short='Café, corte recto.',
        description='Pantalón café de corte recto con pretina de botón, presillas para cinturón y dobladillo amplio.',
        fabric='Gabardina 65 % poliéster · 35 % rayón', care='Lavar en frío; planchar del revés.',
        features=['Rodillas reforzadas', 'Dobladillo de 5 cm para alargar'], featured=True),
    ('Cinturon', 'Niño'): dict(
        id='cl-113', slug='cinturon-nino', line_id='diario', type_id='accesorios', name='Cinturón niño',
        short='Café, hebilla metálica.', embroidery=False,
        description='Cinturón café del uniforme diario para niño.'),
    ('Chamarra invierno', 'Unisex'): dict(
        id='cl-114', slug='chamarra-invierno', line_id='diario', type_id='chamarras', name='Chamarra de invierno',
        short='Chamarra para frío, niña y niño.',
        description='Chamarra de invierno del uniforme, para niña y niño.'),
    ('Sudadera con relieve', 'Unisex'): dict(
        id='cl-115', slug='sudadera-relieve', line_id='diario', type_id='sudaderas', name='Sudadera con relieve',
        short='Sudadera con escudo en relieve.',
        description='Sudadera del uniforme con el escudo en relieve, para niña y niño.'),
    ('Chamarra deportes', 'Unisex'): dict(
        id='cl-116', slug='chamarra-deportes', line_id='deportes', type_id='chamarras', name='Chamarra de deportes',
        short='Marino con franjas naranja y blanca.',
        description='Chamarra deportiva azul marino con cierre al frente, cuello naranja y franjas naranja y blanca en las mangas. Para niña y niño.',
        fabric='Tricot deportivo 100 % poliéster', care='Lavar en frío del revés; no planchar las franjas.',
        features=['Franjas cosidas, no termofijadas', 'Escudo bordado'], featured=True),
    ('Pants de deportes', 'Unisex'): dict(
        id='cl-117', slug='pants-deportes', line_id='deportes', type_id='pants', name='Pants de deportes',
        short='Marino, franjas laterales.',
        description='Pants azul marino con franjas laterales naranja y blanca y cintura de resorte con cordón. Para niña y niño.',
        fabric='Tricot deportivo 100 % poliéster', care='Lavar en frío del revés.',
        features=['Rodillas reforzadas', 'Cordón interior']),
    ('Bata kinder', 'Niña'): dict(
        id='cl-118', slug='bata-kinder-nina', line_id='kinder', type_id='batas', name='Bata kinder niña',
        short='Bata para preescolar.',
        description='Bata del uniforme de kinder para niña.'),
    ('Playera amigos de jesus kinder', 'Unisex'): dict(
        id='cl-119', slug='playera-amigos-de-jesus', line_id='kinder', type_id='playeras', name='Playera Amigos de Jesús',
        short='Blanca con vivos verdes, estampado.',
        description='Playera blanca de kinder con cuello y mangas en verde y el estampado «Amigos de Jesús». Para niña y niño.'),
    ('Chamarra amigos de jesus', 'Unisex'): dict(
        id='cl-120', slug='chamarra-amigos-de-jesus', line_id='kinder', type_id='chamarras', name='Chamarra Amigos de Jesús',
        short='Chamarra de kinder, niña y niño.',
        description='Chamarra «Amigos de Jesús» del uniforme de kinder, para niña y niño.'),
}


def slugify(text):
    text = unicodedata.normalize('NFKD', text).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', text).strip('-')


def photo_for(key, p):
    """Photo named after the Odoo product (colegio + productos + sexo); else the PRODUCTS crop; else sin-foto."""
    path = f"/img/producto/{slugify(' '.join((COLEGIO,) + key))}.webp"
    if (ROOT / path.lstrip('/')).exists():
        return dict(image=path, focus='50% 50%', zoom=1)
    return dict(image=p.get('image', NO_PHOTO), focus=p.get('focus', '50% 50%'), zoom=p.get('zoom', 1))


def size_of(name):
    m = re.search(r'Talla (.+)$', name.strip())
    raw = m.group(1).strip()
    return SIZE.get(raw, raw.lstrip('0') or raw)


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.worksheets[0]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head) if h}

    groups, unknown, skipped = {}, set(), []
    for row in ws.iter_rows(min_row=2):
        v = lambda h: row[col[h]].value
        if not v('folio') or not v('productos'):
            if any(c.value for c in row[:9]):
                skipped.append(f"fila {row[0].row}: {v('nombre')}")
            continue
        key = (v('productos').strip(), v('sexo'))
        if key not in PRODUCTS:
            unknown.add(key)
            continue
        row[col['imagen']].value = SITE + photo_for(key, PRODUCTS[key])['image']
        if str(v('activo')).strip().lower() != 'si':
            continue
        groups.setdefault(key, []).append({'size': size_of(v('nombre')), 'sku': v('folio'), 'price_mxn': round(float(v('precio')))})

    if unknown:
        sys.exit(f'Productos sin definir en PRODUCTS: {sorted(unknown)}')

    old = json.loads(DB.read_text(encoding='utf-8'))
    kept = [dict(p, active=False) for p in old['products'] if p.get('includes')]   # paquetes: not in Odoo yet
    products = []
    for key, p in PRODUCTS.items():
        variants = groups.get(key, [])
        if not variants:
            continue
        products.append({
            'id': p['id'], 'slug': p['slug'], 'sku': variants[0]['sku'][:5], 'line_id': p['line_id'], 'type_id': p['type_id'],
            'fit': FIT[key[1]], 'name': p['name'], 'short': p['short'], 'description': p['description'],
            'price_mxn': min(x['price_mxn'] for x in variants), **photo_for(key, p),
            'sizes': [x['size'] for x in variants], 'variants': variants,
            'fabric': p.get('fabric', ''), 'care': p.get('care', ''), 'features': p.get('features', []),
            'embroidery': p.get('embroidery', True), 'availability': 'existencia', 'active': True, 'featured': p.get('featured', False),
        })

    old.update(
        _schema=old['_schema'] if 'variants [' in old['_schema'] else old['_schema'].replace('sizes,', 'sizes, variants [{size, sku, price_mxn}],'),
        _note='Generado con tools/importar-catalogo.py desde Catalogo.xlsx (Odoo). Precios con IVA por talla; sku de cada variante = folio de Odoo. Varias prendas comparten foto: focus/zoom encuadran la prenda.',
        updated_at=date.today().isoformat(),
        products=products + kept,
    )
    DB.write_text(json.dumps(old, ensure_ascii=False, indent=2) + '\n', encoding='utf-8', newline='\n')
    wb.save(xlsx)

    n = sum(len(p['variants']) for p in products)
    print(f'ok   {DB.relative_to(ROOT)}: {len(products)} productos, {n} variantes ({len(kept)} paquetes desactivados)')
    print(f'ok   {xlsx.name}: columna imagen con la URL de la foto')
    for s in skipped:
        print(f'OJO  {s} (sin folio o sin producto: no se importó)')


if __name__ == '__main__':
    main()
